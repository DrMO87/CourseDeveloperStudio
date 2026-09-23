#!/usr/bin/env python3
"""
Clean Motion Graphic Presentation Deck Generator
Generates academic 16-slide PowerPoint presentations (.pptx) with:
1. 100% Uniform Background (pure white #FFFFFF / high-key modern surface) across ALL slides.
2. Single Unified Sans-Serif Typographic System (Segoe UI / Arial).
3. Simple & Clean Motion Graphic Design Aesthetic:
   - 16:9 Widescreen (13.333" x 7.5")
   - Top Brand Header with University/Faculty logos & Lecture badge
   - Modular Content Cards with subtle borders & clean padding
   - Numbered step pills & key concept cards
   - Embedded authentic figures/diagrams in clean motion-graphic frames
   - Bottom Lecturer Notes drawer
   - Unified footer with slide numbers & accreditation metadata
"""

import os
import sys
import argparse
import re
from typing import List, Dict, Any, Optional

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

try:
    import pptx
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
except ImportError:
    print("Error: python-pptx not installed. Run: pip install python-pptx", file=sys.stderr)
    sys.exit(1)

# Color Palette: Horus University in Egypt (Faculty of Pharmacy)
COLOR_NAVY = RGBColor(0x00, 0x20, 0x60)       # Primary Horus Navy #002060
COLOR_GOLD = RGBColor(0xD9, 0x77, 0x06)       # Warm Accent Amber/Gold #D97706
COLOR_GOLD_LIGHT = RGBColor(0xFE, 0xF3, 0xC7) # Soft Gold Fill #FEF3C7
COLOR_TEXT_MAIN = RGBColor(0x0F, 0x17, 0x2A)  # Slate-900 #0F172A
COLOR_TEXT_MUTED = RGBColor(0x47, 0x55, 0x69) # Slate-600 #475569
COLOR_CARD_BG = RGBColor(0xF8, 0xFA, 0xFC)    # Card Surface Slate-50 #F8FAFC
COLOR_BORDER = RGBColor(0xE2, 0xE8, 0xF0)     # Card Border Slate-200 #E2E8F0
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)      # Slide Canvas Pure White #FFFFFF
COLOR_ACCENT_LINE = RGBColor(0x00, 0x20, 0x60)

FONT_FAMILY = "Segoe UI"


class MotionGraphicDeckBuilder:
    def __init__(
        self,
        course_name: str,
        course_code: str,
        session_code: str,
        session_title: str,
        institution: str,
        assets_dir: Optional[str] = None,
        output_path: Optional[str] = None
    ):
        self.course_name = course_name or "Instrumental Analysis"
        self.course_code = course_code or "PC 206"
        self.session_code = session_code or "Lec 01"
        self.session_title = session_title or "Molecular UV/Vis Spectroscopy"
        self.institution = institution or "Horus University in Egypt (Faculty of Pharmacy)"
        self.assets_dir = assets_dir
        self.output_path = output_path

        self.prs = pptx.Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank_layout = self.prs.slide_layouts[6]

        self.logo_hue = None
        self.logo_pharmacy = None
        self.figures = []

        if self.assets_dir and os.path.exists(self.assets_dir):
            hue_candidate = os.path.join(self.assets_dir, "logo-hue.png")
            if os.path.exists(hue_candidate):
                self.logo_hue = hue_candidate

            pharm_candidate = os.path.join(self.assets_dir, "logo-pharmacy.png")
            if os.path.exists(pharm_candidate):
                self.logo_pharmacy = pharm_candidate

            # Collect figures
            for f in sorted(os.listdir(self.assets_dir)):
                if f.lower().startswith("lec_") and any(f.lower().endswith(ext) for ext in [".png", ".jpg", ".jpeg"]):
                    self.figures.append(os.path.join(self.assets_dir, f))

    def _apply_uniform_background(self, slide):
        """Rule 1: Exact same clean high-key background on every slide."""
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_WHITE

    def _add_header_and_footer(self, slide, slide_num: int, total_slides: int = 16):
        """Consistent Top Header & Bottom Footer on every slide."""
        # Top Accent Line
        top_stripe = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), Inches(13.333), Inches(0.08)
        )
        top_stripe.fill.solid()
        top_stripe.fill.fore_color.rgb = COLOR_NAVY
        top_stripe.line.fill.background()

        # Top Header Area
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(9.5), Inches(0.45))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = f"{self.institution.upper()}  •  {self.course_name.upper()} ({self.course_code})"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_MUTED

        # Lecture Badge (Pill)
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(10.5), Inches(0.2), Inches(1.2), Inches(0.38)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_GOLD_LIGHT
        badge.line.color.rgb = COLOR_GOLD
        badge.line.width = Pt(1)
        btf = badge.text_frame
        btf.margin_top = btf.margin_bottom = btf.margin_left = btf.margin_right = 0
        bp = btf.paragraphs[0]
        bp.text = self.session_code.upper()
        bp.alignment = PP_ALIGN.CENTER
        bp.font.name = FONT_FAMILY
        bp.font.size = Pt(9)
        bp.font.bold = True
        bp.font.color.rgb = COLOR_GOLD

        # Pharmacy Logo Top Right
        if self.logo_pharmacy and os.path.exists(self.logo_pharmacy):
            try:
                slide.shapes.add_picture(self.logo_pharmacy, Inches(11.9), Inches(0.15), width=Inches(0.65))
            except Exception:
                pass

        # Bottom Footer Separator
        footer_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.015)
        )
        footer_line.fill.solid()
        footer_line.fill.fore_color.rgb = COLOR_BORDER
        footer_line.line.fill.background()

        # Footer Text
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(8.0), Inches(0.3))
        ftf = footer_box.text_frame
        ftf.margin_top = ftf.margin_bottom = ftf.margin_left = ftf.margin_right = 0
        fp = ftf.paragraphs[0]
        fp.text = f"{self.session_code}: {self.session_title}  |  Accredited Academic Curriculum  |  HUE Faculty of Pharmacy"
        fp.font.name = FONT_FAMILY
        fp.font.size = Pt(8.5)
        fp.font.color.rgb = COLOR_TEXT_MUTED

        # Slide Number
        page_box = slide.shapes.add_textbox(Inches(10.5), Inches(7.1), Inches(2.0), Inches(0.3))
        ptf = page_box.text_frame
        ptf.margin_top = ptf.margin_bottom = ptf.margin_left = ptf.margin_right = 0
        pp = ptf.paragraphs[0]
        pp.alignment = PP_ALIGN.RIGHT
        pp.text = f"{slide_num:02d} / {total_slides:02d}"
        pp.font.name = FONT_FAMILY
        pp.font.size = Pt(9)
        pp.font.bold = True
        pp.font.color.rgb = COLOR_NAVY

    def _add_slide_title(self, slide, title: str, subtitle: Optional[str] = None):
        """Unified slide title layout with gold accent indicator."""
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.85))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_FAMILY
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.name = FONT_FAMILY
            p2.font.size = Pt(11)
            p2.font.bold = False
            p2.font.color.rgb = COLOR_TEXT_MUTED
            p2.space_before = Pt(3)

        # Subtle gold accent dash
        accent = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), Inches(1.58), Inches(0.8), Inches(0.04)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = COLOR_GOLD
        accent.line.fill.background()

    def create_title_slide(self):
        """Slide 1: Clean, high-impact Title Slide."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._apply_uniform_background(slide)
        self._add_header_and_footer(slide, 1)

        # Hero Motion-Graphic Card Container
        hero_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), Inches(1.75), Inches(11.733), Inches(5.0)
        )
        hero_card.fill.solid()
        hero_card.fill.fore_color.rgb = COLOR_CARD_BG
        hero_card.line.color.rgb = COLOR_BORDER
        hero_card.line.width = Pt(1.5)

        # Left Column: Titles & Details
        details_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(7.5), Inches(4.3))
        tf = details_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0

        p_badge = tf.paragraphs[0]
        p_badge.text = f"ACADEMIC LECTURE SERIES  •  {self.course_code}"
        p_badge.font.name = FONT_FAMILY
        p_badge.font.size = Pt(11)
        p_badge.font.bold = True
        p_badge.font.color.rgb = COLOR_GOLD

        p_title = tf.add_paragraph()
        p_title.text = self.session_title
        p_title.font.name = FONT_FAMILY
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_NAVY
        p_title.space_before = Pt(10)

        p_sub = tf.add_paragraph()
        p_sub.text = f"{self.course_name}  |  {self.session_code}"
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(15)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_TEXT_MAIN
        p_sub.space_before = Pt(6)

        p_desc = tf.add_paragraph()
        p_desc.text = "Comprehensive theoretical principles, optical mechanisms, governing equations, and systematic qualitative/quantitative analytical methodology grounded on original faculty syllabus."
        p_desc.font.name = FONT_FAMILY
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = COLOR_TEXT_MUTED
        p_desc.space_before = Pt(14)

        # Lecturer & Department info card
        info_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(1.2), Inches(5.2), Inches(6.8), Inches(1.2)
        )
        info_card.fill.solid()
        info_card.fill.fore_color.rgb = COLOR_WHITE
        info_card.line.color.rgb = COLOR_BORDER
        info_card.line.width = Pt(1)

        itf = info_card.text_frame
        itf.margin_top = Inches(0.12)
        itf.margin_left = Inches(0.2)
        ip1 = itf.paragraphs[0]
        ip1.text = "Course Coordinator: Dr. Mahmoud Medhat Elkhoudary"
        ip1.font.name = FONT_FAMILY
        ip1.font.size = Pt(11)
        ip1.font.bold = True
        ip1.font.color.rgb = COLOR_NAVY

        ip2 = itf.add_paragraph()
        ip2.text = f"Department of Pharmaceutical Chemistry  •  {self.institution}"
        ip2.font.name = FONT_FAMILY
        ip2.font.size = Pt(9.5)
        ip2.font.color.rgb = COLOR_TEXT_MUTED
        ip2.space_before = Pt(3)

        # Right Column: Visual Frame with University Logo
        if self.logo_hue and os.path.exists(self.logo_hue):
            try:
                slide.shapes.add_picture(self.logo_hue, Inches(9.2), Inches(2.6), width=Inches(2.8))
            except Exception:
                pass

    def create_card_slide(
        self,
        slide_num: int,
        title: str,
        subtitle: str,
        cards: List[Dict[str, Any]],
        image_path: Optional[str] = None,
        lecturer_note: Optional[str] = None
    ):
        """Create a motion-graphic slide with 2 or 3 structured cards and optional image."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._apply_uniform_background(slide)
        self._add_header_and_footer(slide, slide_num)
        self._add_slide_title(slide, title, subtitle)

        content_top = Inches(1.75)
        content_bottom = Inches(6.0) if lecturer_note else Inches(6.85)
        content_height = content_bottom - content_top

        has_image = image_path is not None and os.path.exists(image_path)
        cards_width = Inches(7.3) if has_image else Inches(11.733)

        num_cards = len(cards)
        if num_cards == 1:
            card_height = content_height
            spacing = Inches(0)
        elif num_cards == 2:
            spacing = Inches(0.2)
            card_height = (content_height - spacing) / 2
        else: # 3 cards
            spacing = Inches(0.18)
            card_height = (content_height - (spacing * (num_cards - 1))) / num_cards

        # Render Cards
        for idx, c in enumerate(cards):
            c_top = content_top + idx * (card_height + spacing)
            card_shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.8), c_top, cards_width, card_height
            )
            card_shape.fill.solid()
            card_shape.fill.fore_color.rgb = COLOR_CARD_BG
            card_shape.line.color.rgb = COLOR_BORDER
            card_shape.line.width = Pt(1)

            # Accent Left Indicator Bar
            accent_bar = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0.8), c_top, Inches(0.08), card_height
            )
            accent_bar.fill.solid()
            accent_bar.fill.fore_color.rgb = COLOR_GOLD if idx == 0 else COLOR_NAVY
            accent_bar.line.fill.background()

            # Text inside Card
            ctf = card_shape.text_frame
            ctf.word_wrap = True
            ctf.margin_top = Inches(0.12)
            ctf.margin_left = Inches(0.28)
            ctf.margin_right = Inches(0.2)
            ctf.margin_bottom = Inches(0.1)

            p_h = ctf.paragraphs[0]
            tag = c.get("tag", f"{idx + 1:02d}")
            p_h.text = f"[{tag}]  {c.get('heading', '')}"
            p_h.font.name = FONT_FAMILY
            p_h.font.size = Pt(12)
            p_h.font.bold = True
            p_h.font.color.rgb = COLOR_NAVY

            bullets = c.get("bullets", [])
            for b in bullets:
                p_b = ctf.add_paragraph()
                p_b.text = f"•  {b}"
                p_b.font.name = FONT_FAMILY
                p_b.font.size = Pt(10)
                p_b.font.color.rgb = COLOR_TEXT_MAIN
                p_b.space_before = Pt(2)

        # Render Image on Right Column if present
        if has_image:
            img_x = Inches(8.35)
            img_width = Inches(4.18)
            img_card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                img_x, content_top, img_width, content_height
            )
            img_card.fill.solid()
            img_card.fill.fore_color.rgb = COLOR_WHITE
            img_card.line.color.rgb = COLOR_BORDER
            img_card.line.width = Pt(1)

            try:
                # Add picture centered in the frame
                slide.shapes.add_picture(
                    image_path,
                    img_x + Inches(0.15), content_top + Inches(0.3),
                    width=img_width - Inches(0.3)
                )
            except Exception as ex:
                print(f"Warning: Failed to insert image {image_path}: {ex}", file=sys.stderr)

            # Caption tag at bottom of image frame
            img_cap = slide.shapes.add_textbox(img_x, content_top + content_height - Inches(0.4), img_width, Inches(0.35))
            itf = img_cap.text_frame
            itf.margin_top = itf.margin_left = itf.margin_right = itf.margin_bottom = 0
            ip = itf.paragraphs[0]
            ip.alignment = PP_ALIGN.CENTER
            ip.text = f"FIGURE: {os.path.basename(image_path)}"
            ip.font.name = FONT_FAMILY
            ip.font.size = Pt(8.5)
            ip.font.bold = True
            ip.font.color.rgb = COLOR_TEXT_MUTED

        # Bottom Lecturer Note Drawer
        if lecturer_note:
            note_box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.7)
            )
            note_box.fill.solid()
            note_box.fill.fore_color.rgb = COLOR_WHITE
            note_box.line.color.rgb = COLOR_GOLD_LIGHT
            note_box.line.width = Pt(1)

            ntf = note_box.text_frame
            ntf.margin_top = Inches(0.08)
            ntf.margin_left = Inches(0.2)
            np = ntf.paragraphs[0]
            np.text = f"💡 LECTURER DELIVERY NOTE: {lecturer_note}"
            np.font.name = FONT_FAMILY
            np.font.size = Pt(9.5)
            np.font.color.rgb = COLOR_TEXT_MUTED
            np.font.italic = True

    def build_full_16_slide_deck(self):
        """Construct the complete 16-slide academic curriculum presentation."""
        # 1. Title Slide
        self.create_title_slide()

        # 2. Intended Learning Outcomes (ILOs)
        self.create_card_slide(
            2,
            "Intended Learning Outcomes (ILOs) & Cognitive Map",
            "Structured learning milestones aligned with Bloom's Taxonomy ascent",
            [
                {
                    "tag": "K1 • KNOWLEDGE",
                    "heading": "Foundational Physics & Electromagnetic Radiation",
                    "bullets": [
                        "Define dual nature of EMR: particle nature (photons) and wave nature (λ, ν, ῡ).",
                        "Classify spectroscopic techniques involving an exchange of energy (absorption & emission).",
                        "Distinguish between molecular UV/Vis spectroscopy and atomic spectroscopy."
                    ]
                },
                {
                    "tag": "I1 • INTELLECTUAL",
                    "heading": "Theoretical Derivations & Quantitative Relationships",
                    "bullets": [
                        "Formulate and derive the Beer-Lambert Law: A = ε · b · c = -log(T).",
                        "Analyze chemical, instrumental, and physical causes of deviation from linearity.",
                        "Evaluate quantitative calibration curves for concentration determinations."
                    ]
                },
                {
                    "tag": "P1 • PRACTICAL",
                    "heading": "Spectrophotometric Instrumentation & Quality Control",
                    "bullets": [
                        "Identify the 5 core hardware modules: Source, Monochromator, Sample Cell, Detector, Readout.",
                        "Select appropriate optical materials (quartz vs. glass vs. fused silica cuvettes).",
                        "Execute spectrophotometric assays complying with British Pharmacopoeia (BP) standards."
                    ]
                }
            ],
            lecturer_note="Emphasize that mastery of K1, I1, and P1 is required for both the midterm and the final lab practical."
        )

        # 3. Core Definitions: Molecular UV/Vis Spectroscopy
        fig_01 = self.figures[0] if len(self.figures) > 0 else None
        self.create_card_slide(
            3,
            "Molecular UV/Vis Spectroscopy: Core Principles",
            "Fundamental definition, electromagnetic radiation, and qualitative/quantitative scope",
            [
                {
                    "tag": "DEFINITION",
                    "heading": "Study of Matter-Radiation Interaction",
                    "bullets": [
                        "Interaction between electromagnetic radiation (EMR) and molecules or atoms of matter.",
                        "Qualitative Analysis: Identification of pure substances by characteristic absorption peaks (λmax).",
                        "Quantitative Analysis: Precise determination of analyte concentration in solution."
                    ]
                },
                {
                    "tag": "THEORY",
                    "heading": "Dual Nature of Electromagnetic Radiation",
                    "bullets": [
                        "Wave Nature: Transverse wave composed of mutually perpendicular electric (E) and magnetic (M) fields.",
                        "Particle Nature: Composed of quantized energetic wave packets called photons (E = h · ν = h · c / λ).",
                        "Speed of Light: In a vacuum, radiation propagates at c = 3.00 × 10⁸ m/s."
                    ]
                }
            ],
            image_path=fig_01,
            lecturer_note="Stress that when a photon is absorbed by a sample, it is destroyed, and its discrete energy is acquired by the molecule."
        )

        # 4. Wave Parameters & Governing Equations
        self.create_card_slide(
            4,
            "Wave Parameters & Mathematical Equations",
            "Physical properties governing electromagnetic radiation and photon energy",
            [
                {
                    "tag": "EQUATION 1",
                    "heading": "Wavelength (λ) & Wave Number (ῡ)",
                    "bullets": [
                        "Wavelength (λ): Distance between two successive maxima or minima (Units: nm, Å, µm, m).",
                        "Wave Number (ῡ): Reciprocal of wavelength representing waves per centimeter: ῡ = 1 / λ (Unit: cm⁻¹).",
                        "Directly proportional to energy and frequency (ῡ = ν / c)."
                    ]
                },
                {
                    "tag": "EQUATION 2",
                    "heading": "Frequency (ν) & Quantum Energy (E)",
                    "bullets": [
                        "Frequency (ν): Number of complete cycles per second (Unit: s⁻¹, Hz). ν = c / λ.",
                        "Planck Equation: E = h · ν = (h · c) / λ, where h = 6.626 × 10⁻³⁴ J·s.",
                        "Crucial relationship: Higher frequency / shorter wavelength corresponds directly to HIGHER energy."
                    ]
                },
                {
                    "tag": "LIGHT TYPES",
                    "heading": "Polychromatic vs. Monochromatic Radiation",
                    "bullets": [
                        "Polychromatic Light: Beam containing multiple wavelengths (e.g. tungsten/deuterium source output).",
                        "Monochromatic Light: Beam of a single, isolated discrete wavelength (required for Beer's Law validity)."
                    ]
                }
            ],
            lecturer_note="Make sure students understand why wavenumber (cm⁻¹) is favored in infrared while nanometer (nm) is used in UV-Vis."
        )

        # 5. The Electromagnetic Spectrum
        fig_02 = self.figures[1] if len(self.figures) > 1 else None
        self.create_card_slide(
            5,
            "The Electromagnetic Spectrum & Optical Bands",
            "Spectral division, transition types, and pharmaceutical analytical ranges",
            [
                {
                    "tag": "SPECTRAL RANGE",
                    "heading": "UV-Visible Analytical Regions",
                    "bullets": [
                        "Vacuum Ultraviolet: 10 - 200 nm (absorbed by atmospheric oxygen; requires vacuum apparatus).",
                        "Near UV Region: 200 - 400 nm (quartz cuvettes required; primary region for aromatic pharmaceuticals).",
                        "Visible Region: 400 - 800 nm (glass/plastic cuvettes permitted; colored complexes & chromophores)."
                    ]
                },
                {
                    "tag": "ENERGY IMPACT",
                    "heading": "Molecular Energy State Transitions",
                    "bullets": [
                        "UV-Vis Absorption: Promotes valence electrons from ground state to higher molecular orbitals (electronic transition).",
                        "Infrared (IR): Induces vibrational transitions between atomic bonds.",
                        "Microwave / Radiowaves: Induces rotational states and nuclear spin flips (NMR)."
                    ]
                }
            ],
            image_path=fig_02,
            lecturer_note="Remind students that quartz is transparent down to 190 nm, whereas standard optical glass absorbs strongly below 350 nm."
        )

        # 6. Classification of Spectroscopic Techniques
        self.create_card_slide(
            6,
            "Classification of Spectroscopic Techniques",
            "Energy exchange modalities: absorption, emission, luminescence, and scattering",
            [
                {
                    "tag": "CLASS 1",
                    "heading": "Absorption Spectroscopy",
                    "bullets": [
                        "Incident radiation is absorbed by analyte ground-state molecules: M + hν → M*.",
                        "Techniques: UV-Vis Spectrophotometry, Infrared (IR) Spectroscopy, Atomic Absorption (AAS).",
                        "Directly quantifies light attenuation through the sample medium."
                    ]
                },
                {
                    "tag": "CLASS 2",
                    "heading": "Emission & Luminescence Spectroscopy",
                    "bullets": [
                        "Analyte molecules are excited by thermal/electrical/optical energy and emit photons upon relaxation: M* → M + hν.",
                        "Fluorescence & Phosphorescence: Photoluminescence re-emitting light at longer wavelengths (Stokes shift).",
                        "Flame Photometry / Atomic Emission: Thermal excitation for alkali/alkaline earth metals (Na, K, Li, Ca)."
                    ]
                },
                {
                    "tag": "CLASS 3",
                    "heading": "Scattering & Refraction Methods",
                    "bullets": [
                        "Turbidimetry & Nephelometry: Measurement of light scattered by un-dissolved particulate matter in suspension.",
                        "Refractometry: Measurement of the angle of refraction reflecting refractive index (η)."
                    ]
                }
            ],
            lecturer_note="Emphasize that absorption measures the decrease in transmitted beam intensity, whereas emission measures emitted intensity against a zero baseline."
        )

        # 7. Energy States of Molecules & Orbital Transitions
        fig_03 = self.figures[2] if len(self.figures) > 2 else None
        self.create_card_slide(
            7,
            "Molecular Electronic Transitions & Energy States",
            "Total molecular energy components: electronic, vibrational, and rotational",
            [
                {
                    "tag": "ENERGY EQUATION",
                    "heading": "Total Internal Molecular Energy (Etotal)",
                    "bullets": [
                        "Etotal = Eelectronic + Evibrational + Erotational.",
                        "Energy quantization: Eelectronic (1-10 eV) >> Evibrational (0.1 eV) >> Erotational (0.001 eV).",
                        "Electronic excitation produces complex broad bands rather than sharp lines due to overlaid vibrational sub-levels."
                    ]
                },
                {
                    "tag": "ORBITALS",
                    "heading": "Types of Molecular Orbitals Involved",
                    "bullets": [
                        "σ (Bonding) & σ* (Antibonding): High-energy single bonds; require vacuum UV (< 180 nm).",
                        "π (Bonding) & π* (Antibonding): Unsaturated double/triple bonds (chromophores); 200 - 700 nm.",
                        "n (Non-bonding): Unshared electron pairs on heteroatoms (O, N, S, Halogens); n → π* transitions."
                    ]
                }
            ],
            image_path=fig_03,
            lecturer_note="Point out why UV spectra of molecules in solution appear as smooth, broad absorption bands rather than discrete atomic lines."
        )

        # 8. Systematic Instrumentation: The 5 Core Modules
        fig_04 = self.figures[3] if len(self.figures) > 3 else None
        self.create_card_slide(
            8,
            "Spectrophotometer Architecture: 5 Essential Modules",
            "Optical path and hardware sequence from radiation source to electronic readout",
            [
                {
                    "tag": "MODULES 1 & 2",
                    "heading": "1. Radiation Source & 2. Monochromator",
                    "bullets": [
                        "Source: Deuterium discharge lamp (UV: 190-380 nm) + Tungsten-Halogen filament (Visible: 350-900 nm).",
                        "Monochromator: Entrance slit → Collimator → Dispersing element (Prism or Holographic Diffraction Grating) → Exit slit.",
                        "Function: Isolates a narrow band of nominal wavelength (effective bandwidth)."
                    ]
                },
                {
                    "tag": "MODULES 3 - 5",
                    "heading": "3. Cuvette, 4. Detector & 5. Readout",
                    "bullets": [
                        "Sample Cell: Precision quartz (UV/Vis) or optical glass (Vis only); standard path length b = 1.00 cm.",
                        "Photodetector: Photodiode Array (PDA) or Photomultiplier Tube (PMT) converting photons into proportional electric current.",
                        "Readout: Analog-to-digital converter feeding microprocessor for absorbance computation."
                    ]
                }
            ],
            image_path=fig_04,
            lecturer_note="Draw the sequential block diagram on the board: Source → Monochromator → Sample → Detector → Readout."
        )

        # 9. Optical Monochromators: Prisms vs. Diffraction Gratings
        self.create_card_slide(
            9,
            "Optical Dispersing Elements: Prisms vs. Gratings",
            "Comparative physics, dispersion geometry, and resolving power",
            [
                {
                    "tag": "ELEMENT 1",
                    "heading": "Refraction Prisms (Quartz / Glass)",
                    "bullets": [
                        "Mechanism: Operates via Snell's Law and wavelength-dependent refractive index (dispersion).",
                        "Non-linear Dispersion: Shorter wavelengths (blue/UV) are refracted much more sharply than longer wavelengths (red).",
                        "Limitation: Crowded dispersion at longer wavelengths; temperature-sensitive."
                    ]
                },
                {
                    "tag": "ELEMENT 2",
                    "heading": "Diffraction Gratings (Holographic)",
                    "bullets": [
                        "Mechanism: Polished surface ruled with thousands of closely spaced parallel grooves (e.g. 1200 lines/mm).",
                        "Linear Dispersion: Produces constant linear dispersion across the entire UV-Vis spectrum: d · sin(θ) = m · λ.",
                        "Advantages: Superior optical resolution, uniform dispersion, and eliminates prism absorption losses."
                    ]
                }
            ],
            lecturer_note="Explain why modern research-grade UV-Vis spectrophotometers universally utilize holographic diffraction gratings."
        )

        # 10. Single-Beam vs. Double-Beam Optical Configurations
        fig_05 = self.figures[4] if len(self.figures) > 4 else None
        self.create_card_slide(
            10,
            "Single-Beam vs. Double-Beam Spectrophotometers",
            "Beam splitting, drift compensation, and baseline stability",
            [
                {
                    "tag": "SINGLE-BEAM",
                    "heading": "Single-Beam Architecture",
                    "bullets": [
                        "Single optical path: User must alternate between blank and sample cuvettes at each wavelength.",
                        "Susceptible to source intensity fluctuations, detector drift, and voltage variations.",
                        "Best suited for routine quantitative analysis at a fixed single wavelength."
                    ]
                },
                {
                    "tag": "DOUBLE-BEAM",
                    "heading": "Double-Beam in Space / Time",
                    "bullets": [
                        "Optical beam chopper splits beam into two paths: Sample Beam (I) and Reference/Blank Beam (I₀).",
                        "Continuous real-time ratioing (I / I₀) automatically cancels source drift, solvent absorption, and line voltage spikes.",
                        "Required for automatic spectral scanning across a broad wavelength range (200 - 800 nm)."
                    ]
                }
            ],
            image_path=fig_05,
            lecturer_note="Highlight that double-beam spectrophotometers eliminate the need to manually re-zero the blank when scanning spectra."
        )

        # 11. Photodetectors: Phototubes, PMTs & Photodiode Arrays
        fig_06 = self.figures[5] if len(self.figures) > 5 else None
        self.create_card_slide(
            11,
            "Radiation Detectors: PMTs & Photodiode Arrays",
            "Photoelectric effect, electron multiplication, and simultaneous multi-channel detection",
            [
                {
                    "tag": "PMT",
                    "heading": "Photomultiplier Tube (PMT)",
                    "bullets": [
                        "Photoemissive cathode emits electrons upon photon impact (photoelectric effect).",
                        "Series of dynodes at successive higher voltages causes secondary electron cascade (amplification 10⁶ - 10⁷).",
                        "Extremely sensitive; gold standard for low light levels and ultra-trace quantification."
                    ]
                },
                {
                    "tag": "PDA / CCD",
                    "heading": "Photodiode Array (PDA) Detector",
                    "bullets": [
                        "Linear array of hundreds of microscopic silicon photodiodes on a single semiconductor chip.",
                        "Polychromatic light passes through sample first; then dispersed across the array (reverse optics).",
                        "Simultaneous detection: Captures entire UV-Vis spectrum instantaneously (milliseconds); essential for HPLC UV-detectors."
                    ]
                }
            ],
            image_path=fig_06,
            lecturer_note="Point out that PDA detectors enable 3D contour spectral plots in pharmaceutical chromatographic testing."
        )

        # 12. Quantitative Derivations: Beer-Lambert Law
        fig_07 = self.figures[6] if len(self.figures) > 6 else None
        self.create_card_slide(
            12,
            "Quantitative Laws of Absorption: Beer-Lambert",
            "Fundamental mathematical derivations governing radiant energy attenuation",
            [
                {
                    "tag": "LAMBERT'S LAW",
                    "heading": "Path Length Proportionality (Bouguer-Lambert)",
                    "bullets": [
                        "Absorbance is directly proportional to the thickness/path length of the absorbing medium: A ∝ b.",
                        "Each successive layer of equal thickness absorbs an equal fraction of radiant energy traversing it: -dI / db = k₁ · I."
                    ]
                },
                {
                    "tag": "BEER'S LAW",
                    "heading": "Concentration Proportionality (Beer)",
                    "bullets": [
                        "Absorbance is directly proportional to the concentration of the absorbing chemical species: A ∝ c.",
                        "Combined Formula: A = ε · b · c, where ε is the Molar Absorptivity (L·mol⁻¹·cm⁻¹).",
                        "Transmittance Relation: T = I / I₀; %T = (I / I₀) × 100; Absorbance A = -log₁₀(T) = 2 - log₁₀(%T)."
                    ]
                }
            ],
            image_path=fig_07,
            lecturer_note="Demonstrate that while Transmittance decreases exponentially with concentration, Absorbance increases linearly."
        )

        # 13. Molar Absorptivity & Specific Absorbance (A 1%, 1cm)
        self.create_card_slide(
            13,
            "Constants of Absorption: ε vs. A(1%, 1cm)",
            "Pharmacopoeial expression of absorption intensity and quantitative conversion",
            [
                {
                    "tag": "CONSTANT 1",
                    "heading": "Molar Absorptivity (ε)",
                    "bullets": [
                        "Absorbance of a 1.00 M solution of analyte contained in a 1.00 cm cuvette (Unit: L·mol⁻¹·cm⁻¹).",
                        "Characteristic physical constant for a chemical compound at a designated wavelength, solvent, and temperature.",
                        "Magnitude indicates transition probability: ε > 10,000 indicates strongly allowed π → π* transitions."
                    ]
                },
                {
                    "tag": "CONSTANT 2",
                    "heading": "Specific Absorbance: A(1%, 1cm)",
                    "bullets": [
                        "Absorbance of a 1% w/v solution (1 g / 100 mL) in a 1.00 cm cell path length.",
                        "Universally utilized in British Pharmacopoeia (BP) monographs when molecular weight is unknown or complex.",
                        "Conversion Formula: ε = [A(1%, 1cm) × Molecular Weight] / 10."
                    ]
                }
            ],
            lecturer_note="Have students practice converting between A(1%, 1cm) and molar absorptivity ε using an active pharmaceutical ingredient example."
        )

        # 14. Deviations from Beer-Lambert Law
        fig_08 = self.figures[7] if len(self.figures) > 7 else None
        self.create_card_slide(
            14,
            "Deviations from the Beer-Lambert Law",
            "Systematic chemical, instrumental, and physical causes of non-linearity",
            [
                {
                    "tag": "CAUSE 1",
                    "heading": "Real / Chemical Deviations",
                    "bullets": [
                        "High Concentration (> 0.01 M): Electrostatic interactions between neighboring ions alter electron charge distribution.",
                        "Chemical Equilibria: Solute-solvent association, dissociation, polymerization, or tautomerization (e.g. pH-sensitive dyes).",
                        "Refractive Index Changes: High solute concentrations shift the medium's refractive index (η)."
                    ]
                },
                {
                    "tag": "CAUSE 2",
                    "heading": "Instrumental & Optical Deviations",
                    "bullets": [
                        "Polychromatic Radiation: Beer's law strictly applies to monochromatic light; wide slit widths cause negative deviations.",
                        "Stray Light: Unwanted radiation of foreign wavelengths reaching detector; places a severe ceiling on maximum measurable absorbance.",
                        "Mismatched Cuvettes: Differences in optical path or dirty optical faces generate baseline offsets."
                    ]
                }
            ],
            image_path=fig_08,
            lecturer_note="Stress that ideal analytical absorbance measurements should always fall within the linear range of 0.2 to 0.8 absorbance units."
        )

        # 15. Summary Matrix & Conceptual Synthesis
        self.create_card_slide(
            15,
            "Conceptual Synthesis & Method Comparison",
            "Unified summary matrix of spectroscopic principles and analytical decision rules",
            [
                {
                    "tag": "PHYSICS SUMMARY",
                    "heading": "Governing Physical Laws",
                    "bullets": [
                        "E = h · ν = (h · c) / λ  |  Wavenumber ῡ = 1 / λ (cm⁻¹).",
                        "Absorption occurs when photon energy exactly matches the electronic energy gap (ΔE = E₂ - E₁).",
                        "UV (200-400 nm) requires quartz; Visible (400-800 nm) allows optical glass."
                    ]
                },
                {
                    "tag": "INSTRUMENT SUMMARY",
                    "heading": "Hardware Best Practices",
                    "bullets": [
                        "Double-beam instrumentation cancels optical fluctuations and drift.",
                        "Holographic gratings ensure linear dispersion across all analytical wavelengths.",
                        "PDA detectors allow instantaneous multi-wavelength UV detection."
                    ]
                },
                {
                    "tag": "MATH SUMMARY",
                    "heading": "Quantitative Analysis Rules",
                    "bullets": [
                        "A = ε · b · c  |  A = -log(T) = 2 - log(%T).",
                        "Concentrations must remain dilute (< 0.01 M) to avoid electrostatic deviation.",
                        "Target absorbance range for minimum photometric error: 0.2 - 0.8 AU."
                    ]
                }
            ],
            lecturer_note="Review this summary matrix before ending the lecture to anchor the key takeaways."
        )

        # 16. Self-Assessment Checkpoints & Review Problems
        self.create_card_slide(
            16,
            "Self-Assessment Checkpoints & Exam Preparation",
            "Mastery verification questions directly aligned with lecture ILOs",
            [
                {
                    "tag": "QUESTION 1",
                    "heading": "Derivation & Calculation (Bloom: Apply)",
                    "bullets": [
                        "A solution of a pharmaceutical analyte (MW = 250 g/mol) exhibits A = 0.600 in a 1.00 cm cuvette.",
                        "If the molar absorptivity ε = 15,000 L·mol⁻¹·cm⁻¹, calculate: (a) Molar concentration, (b) Concentration in µg/mL.",
                        "Calculate the Transmittance (T) and %T of this sample."
                    ]
                },
                {
                    "tag": "QUESTION 2",
                    "heading": "Theoretical Distinction (Bloom: Understand)",
                    "bullets": [
                        "Explain why quartz cuvettes are mandatory for UV measurements while glass cuvettes can be used in the visible range.",
                        "Why does an increase in stray light cause a NEGATIVE deviation from Beer's Law at high concentrations?"
                    ]
                },
                {
                    "tag": "QUESTION 3",
                    "heading": "Practical Diagnostic (Bloom: Evaluate)",
                    "bullets": [
                        "A student attempts to measure a 0.1 M drug solution and obtains a non-linear calibration curve. Detail two physical reasons for this deviation."
                    ]
                }
            ],
            lecturer_note="Assign Question 1 as a mandatory formative checkpoint for the next laboratory practical session."
        )

    def save(self, path: Optional[str] = None):
        target = path or self.output_path or "MotionGraphic_Deck.pptx"
        os.makedirs(os.path.dirname(os.path.abspath(target)), exist_ok=True)
        self.prs.save(target)
        print(f"✅ Motion Graphic PPTX Deck created successfully: {target}")
        return target


def main():
    parser = argparse.ArgumentParser(description="Generate Academic Motion Graphic PowerPoint Presentation (.pptx)")
    parser.add_argument("--project-slug", default="inst", help="Course Project Slug")
    parser.add_argument("--session-code", default="Lec 01", help="Lecture session code (e.g. Lec 01)")
    parser.add_argument("--session-title", default="Principles of Molecular UV/Vis Spectroscopy", help="Lecture Title")
    parser.add_argument("--course-name", default="Instrumental Analysis", help="Course Name")
    parser.add_argument("--course-code", default="PC 206", help="Course Code")
    parser.add_argument("--institution", default="Horus University in Egypt (Faculty of Pharmacy)", help="Institution Name")
    parser.add_argument("--assets-dir", default=None, help="Directory containing extracted figures & logos")
    parser.add_argument("--output-path", default=None, help="Destination PPTX file path")
    args = parser.parse_args()

    # Determine default paths if not specified
    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    
    assets_dir = args.assets_dir
    if not assets_dir:
        candidates = [
            os.path.join(base_dir, "vaults", "Inst-Analysis", "01_Projects", "Inst-Analysis", args.session_code, "_assets"),
            os.path.join(base_dir, "vaults", "Inst-Analysis", "01_Projects", "Inst-Analysis", args.session_code.replace(" ", "_"), "_assets"),
            os.path.join(base_dir, "vaults", args.project_slug, "01_Projects", args.project_slug, args.session_code, "_assets"),
            os.path.join(base_dir, "vaults", "Inst-Analysis", "01_Projects", "Inst-Analysis", "Dossier", "_assets"),
        ]
        for c in candidates:
            if os.path.exists(c):
                assets_dir = c
                break

    output_path = args.output_path
    if not output_path:
        clean_code = args.session_code.replace(" ", "_")
        output_path = os.path.join(
            base_dir, "vaults", "Inst-Analysis", "03_Resources", "NotebookLM_Generated",
            f"{args.course_name.replace(' ', '_')}_{clean_code}_MotionGraphic_Deck.pptx"
        )

    builder = MotionGraphicDeckBuilder(
        course_name=args.course_name,
        course_code=args.course_code,
        session_code=args.session_code,
        session_title=args.session_title,
        institution=args.institution,
        assets_dir=assets_dir,
        output_path=output_path
    )

    builder.build_full_16_slide_deck()
    saved = builder.save(output_path)

    # Also save a copy to .nlm-downloads if present
    dl_staging = os.path.join(base_dir, ".nlm-downloads", "inst")
    if os.path.exists(dl_staging):
        for sub in os.listdir(dl_staging):
            sub_p = os.path.join(dl_staging, sub)
            if os.path.isdir(sub_p):
                for inner in os.listdir(sub_p):
                    inner_p = os.path.join(sub_p, inner)
                    if os.path.isdir(inner_p):
                        try:
                            copy_target = os.path.join(inner_p, os.path.basename(output_path))
                            import shutil
                            shutil.copy2(output_path, copy_target)
                            print(f"✅ Staged copy in: {copy_target}")
                        except Exception:
                            pass


if __name__ == "__main__":
    main()
